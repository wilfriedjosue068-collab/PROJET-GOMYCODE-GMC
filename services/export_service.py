"""UC9 - Export des transactions en CSV ou Excel."""
import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services.transaction_service import compute_totals, list_transactions
from utils.formatters import format_date, type_label

HEADERS = ["Date", "Type", "Catégorie", "Description", "Montant (FCFA)"]


def _rows(transactions):
    for t in transactions:
        amount = t["amount"] if t["type"] == "revenu" else -t["amount"]
        yield [format_date(t["date"]), type_label(t["type"]), t["category"],
               t["description"], amount]


def export_csv(user_id, path, start=None, end=None):
    """Exporte en CSV (séparateur ';', lisible directement dans Excel). Retourne le nb de lignes."""
    transactions = list_transactions(user_id, start=start, end=end)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(HEADERS)
        writer.writerows(_rows(transactions))
    return len(transactions)


def export_excel(user_id, path, start=None, end=None):
    """Exporte en .xlsx avec en-têtes stylés et totaux. Retourne le nb de lignes."""
    transactions = list_transactions(user_id, start=start, end=end)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E7C66")
        cell.alignment = Alignment(horizontal="center")
    for row in _rows(transactions):
        ws.append(row)
        ws.cell(ws.max_row, 5).number_format = '#,##0;[Red]-#,##0'

    totals = compute_totals(transactions)
    ws.append([])
    for label, value in (("Total revenus", totals["revenus"]),
                         ("Total dépenses", -totals["depenses"]),
                         ("Solde", totals["solde"])):
        ws.append(["", "", "", label, value])
        ws.cell(ws.max_row, 4).font = Font(bold=True)
        ws.cell(ws.max_row, 5).font = Font(bold=True)
        ws.cell(ws.max_row, 5).number_format = '#,##0;[Red]-#,##0'

    for col, width in zip("ABCDE", (12, 10, 22, 40, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(transactions)
