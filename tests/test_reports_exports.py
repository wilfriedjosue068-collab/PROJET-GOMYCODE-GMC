import csv

from openpyxl import load_workbook

from services import export_service, report_service
from services import transaction_service as ts
from utils.dates import month_label, shift_month
from utils.formatters import format_amount


def _seed(user_id, cats):
    ts.add_transaction(user_id, cats["Salaire"], 450000, "28/09/2026")
    ts.add_transaction(user_id, cats["Transport"], 20000, "02/09/2026")
    ts.add_transaction(user_id, cats["Alimentation"], 60000, "05/09/2026")
    ts.add_transaction(user_id, cats["Salaire"], 450000, "28/08/2026")
    ts.add_transaction(user_id, cats["Logement"], 120000, "05/08/2026")


def test_monthly_summary(user_id, cats):
    _seed(user_id, cats)
    s = report_service.get_monthly_summary(user_id, "2026-09")
    assert (s["revenus"], s["depenses"], s["solde"]) == (450000, 80000, 370000)
    assert report_service.get_global_balance(user_id) == 700000


def test_expenses_by_category(user_id, cats):
    _seed(user_id, cats)
    assert report_service.get_expenses_by_category(user_id, "2026-09") == [
        ("Alimentation", 60000), ("Transport", 20000)]


def test_monthly_trend(user_id, cats):
    _seed(user_id, cats)
    trend = report_service.get_monthly_trend(user_id, "2026-09", months=3)
    assert [t["month"] for t in trend] == ["2026-07", "2026-08", "2026-09"]
    assert trend[1]["depenses"] == 120000 and trend[0]["revenus"] == 0


def test_exports(user_id, cats, tmp_path):
    _seed(user_id, cats)
    csv_path = tmp_path / "export.csv"
    assert export_service.export_csv(user_id, csv_path, start="01/09/2026") == 3
    with open(csv_path, encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[0][0] == "Date" and len(rows) == 4

    xlsx_path = tmp_path / "export.xlsx"
    assert export_service.export_excel(user_id, xlsx_path) == 5
    ws = load_workbook(xlsx_path).active
    assert ws["A1"].value == "Date"
    assert ws.cell(ws.max_row, 5).value == 700000   # ligne "Solde"


def test_formatters_and_dates():
    assert format_amount(137500) == "137 500 FCFA"
    assert format_amount(-2500, with_currency=False) == "-2 500"
    assert shift_month("2026-01", -1) == "2025-12"
    assert shift_month("2026-11", 3) == "2027-02"
    assert month_label("2026-09") == "Septembre 2026"
